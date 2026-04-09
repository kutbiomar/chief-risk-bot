from __future__ import annotations

import io

from fastapi import status
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from backend.models.analytics import VarResult
from backend.models.content import Document, ExtractionArtifact, ExtractionResult
from backend.models.content import BriefingRun
from backend.models.portfolio import PortfolioSnapshot, Position
from backend.services.auth.session import utc_now
from backend.services.briefings import PdfExportUnavailableError
from backend.tests.test_auth import seed_user


def bootstrap_portfolio(client: TestClient, db_session: Session, email: str = "phasecd@example.com") -> dict[str, str]:
    user = seed_user(db_session, email=email)
    login = client.post("/api/auth/login", json={"email": user.email, "password": "secret123"})
    assert login.status_code == 200
    csrf = login.cookies.get("__crb_csrf", "")
    payload = (
        "ticker,quantity,asset_class,custodian,geo_region,sector,market_segment,market_value_usd\n"
        "AAPL,10,public_equity,Goldman,US,Technology,Large Cap,1900\n"
        "BND,20,fixed_income,Fidelity,US,Fixed Income,IG Credit,1500\n"
        "VGK,15,public_equity,Goldman,Europe,Financials,Large Cap,800\n"
        "PEF,5,private_equity,AltCustodian,Global,Alternatives,Private Equity,600\n"
    )
    upload = client.post(
        "/api/ingest/csv",
        files={"file": ("positions.csv", payload, "text/csv")},
        headers={"X-CSRF-Token": csrf},
    )
    assert upload.status_code == 200
    return {"csrf": csrf, "snapshot_id": upload.json()["snapshot_id"]}


def csrf_headers(auth: dict[str, str]) -> dict[str, str]:
    return {"X-CSRF-Token": auth["csrf"]}


def build_statement_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Holdings"
    sheet.append(
        [
            "Ticker",
            "Security Name",
            "Quantity",
            "Market Value USD",
            "Asset Class",
            "Custodian",
            "Region",
            "Sector",
            "Market Segment",
        ]
    )
    sheet.append(["MSFT", "Microsoft Corp", 12, 5040, "Public Equity", "Goldman", "US", "Technology", "Large Cap"])
    sheet.append(["BND", "Vanguard Total Bond", 40, 3200, "Fixed Income", "Fidelity", "US", "Fixed Income", "IG Credit"])
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def build_capital_call_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Capital Call"
    sheet.append(["Capital Call Notice"])
    sheet.append(["Due Date", "2026-04-30"])
    sheet.append(["Amount", "$1,250,000"])
    sheet.append(["Bank", "First Private Bank"])
    sheet.append(["Account Number", "99887766"])
    sheet.append(["ABA", "021000021"])
    payload = io.BytesIO()
    workbook.save(payload)
    return payload.getvalue()


def test_var_risk_and_cockpit_flow(client: TestClient, db_session: Session) -> None:
    auth = bootstrap_portfolio(client, db_session)

    var_response = client.post("/api/var/compute", headers=csrf_headers(auth))
    assert var_response.status_code == 200
    assert var_response.json()["methodology"] == "historical_simulation"
    assert var_response.json()["effective_lookback_days"] > 200

    risk_response = client.post("/api/risk/run", headers=csrf_headers(auth))
    assert risk_response.status_code == 200
    assert len(risk_response.json()["scores"]) == 5

    cockpit = client.get("/api/cockpit")
    assert cockpit.status_code == 200
    body = cockpit.json()
    assert body["portfolio_summary"]["total_aum_usd"] == 4800.0
    assert body["var_result"]["var_1d_95"] >= 0.0
    assert len(body["risk_register"]) >= 5


def test_briefings_settings_api_keys_and_documents_flow(client: TestClient, db_session: Session) -> None:
    auth = bootstrap_portfolio(client, db_session, email="phasecd2@example.com")
    headers = csrf_headers(auth)
    assert client.post("/api/var/compute", headers=headers).status_code == 200
    assert client.post("/api/risk/run", headers=headers).status_code == 200

    settings_patch = client.patch("/api/settings", json={"briefing_day": "Friday", "briefing_send_pdf": True}, headers=headers)
    assert settings_patch.status_code == 200
    assert settings_patch.json()["briefing_day"] == "Friday"
    assert settings_patch.json()["briefing_send_pdf"] is True

    key_create = client.post("/api/settings/api-keys", json={"label": "Demo", "key_type": "anthropic"}, headers=headers)
    assert key_create.status_code == 200
    assert key_create.json()["plain_text_key"].startswith("crb_")
    key_id = key_create.json()["id"]
    key_list = client.get("/api/settings/api-keys")
    assert key_list.status_code == 200
    assert len(key_list.json()) == 1

    briefing = client.post("/api/briefings/generate", headers=headers)
    assert briefing.status_code == 200
    briefing_id = briefing.json()["id"]
    output = briefing.json()["output"]
    assert "executive_summary" in output
    assert "portfolio_risks" in output
    assert "recommendations" in output
    assert client.post(f"/api/briefings/{briefing_id}/publish", headers=headers).status_code == 200
    export = client.get(f"/api/briefings/{briefing_id}/export/pdf")
    assert export.status_code in {200, status.HTTP_503_SERVICE_UNAVAILABLE}
    if export.status_code == 200:
        assert "attachment;" in export.headers["content-disposition"]
        assert f"week-15-2026_v1" in export.headers["content-disposition"]
        assert export.headers["content-type"].startswith("application/pdf")
        assert len(export.content) > 0
    else:
        assert export.json()["detail"] == "PDF export unavailable — WeasyPrint system libraries not installed"

    document = client.post(
        "/api/documents/upload",
        data={"folder": "custodian_statements"},
        files={
            "file": (
                "statement.xlsx",
                build_statement_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert document.status_code == 200
    document_id = document.json()["id"]
    assert client.post(f"/api/documents/{document_id}/parse", headers=headers).status_code == 200
    extraction = client.get(f"/api/documents/{document_id}/extraction")
    assert extraction.status_code == 200
    extraction_body = extraction.json()
    assert extraction_body["needs_review_count"] == 0
    assert extraction_body["positions"][0]["ticker"] == "MSFT"
    assert extraction_body["positions"][0]["market_value_usd"] == 5040.0
    assert extraction_body["positions"][0]["factor_asset_class"] == "public_equity"
    assert extraction_body["positions"][0]["factor_sector"] == "technology"
    assert extraction_body["positions"][0]["factor_region"] == "us"
    assert extraction_body["positions"][0]["factor_tag_source"] == "extracted"
    assert extraction_body["positions"][0]["factor_tag_confidence"] == 0.82

    parsed_document = db_session.get(Document, document_id)
    assert parsed_document is not None
    assert parsed_document.tag == "nav_statement"
    assert parsed_document.extraction_status == "done"

    extraction_row = db_session.get(ExtractionResult, parsed_document.extraction_result_id)
    assert extraction_row is not None
    assert extraction_row.confidence_json.startswith("[")
    artifacts = db_session.scalars(
        select(ExtractionArtifact).where(ExtractionArtifact.extraction_result_id == extraction_row.id)
    ).all()
    artifact_types = {artifact.artifact_type for artifact in artifacts}
    assert {"layout", "classification", "risk", "treasury", "reconciliation"} <= artifact_types

    review = client.get(f"/api/documents/{document_id}/review")
    assert review.status_code == 200
    review_body = review.json()
    assert review_body["layout"]["parser"] == "xlsx-layout-parser"
    assert review_body["layout"]["row_count"] >= 2
    assert review_body["positions"][0]["source_ref"]["sheet_name"] == "Holdings"

    approve = client.post(f"/api/documents/{document_id}/approve", headers=headers)
    assert approve.status_code == 200
    portfolio_positions = client.get("/api/portfolio/positions")
    assert portfolio_positions.status_code == 200
    assert portfolio_positions.json()["total"] == 2
    assert portfolio_positions.json()["items"][0]["ticker"] == "MSFT"
    assert portfolio_positions.json()["items"][0]["factor_asset_class"] == "public_equity"
    assert portfolio_positions.json()["items"][0]["factor_sector"] == "technology"
    assert portfolio_positions.json()["items"][0]["factor_tag_source"] == "extracted"
    assert portfolio_positions.json()["items"][0]["factor_tag_confidence"] == 0.82

    msft_position = db_session.scalar(select(Position).where(Position.ticker == "MSFT"))
    assert msft_position is not None
    assert msft_position.factor_asset_class == "public_equity"
    assert msft_position.factor_sector == "technology"
    assert msft_position.factor_region == "us"
    assert msft_position.factor_tag_source == "extracted"
    assert msft_position.factor_tag_confidence == 0.82

    assert client.post(f"/api/documents/{document_id}/tag", json={"tag": "imported"}, headers=headers).status_code == 200
    docs = client.get("/api/documents")
    assert docs.status_code == 200
    assert docs.json()["folder_counts"]["custodian_statements"] == 1
    assert client.delete(f"/api/documents/{document_id}", headers=headers).status_code == 200

    broken_pdf = client.post(
        "/api/documents/upload",
        data={"folder": "custodian_statements"},
        files={"file": ("broken.pdf", b"%PDF-1.4\nnot really a pdf", "application/pdf")},
        headers=headers,
    )
    assert broken_pdf.status_code == 200
    broken_parse = client.post(f"/api/documents/{broken_pdf.json()['id']}/parse", headers=headers)
    assert broken_parse.status_code == 400
    assert broken_parse.json()["detail"] == (
        "Unable to parse this PDF - it may be corrupted, password-protected, or not a valid PDF."
    )

    revoke = client.delete(f"/api/settings/api-keys/{key_id}", headers=headers)
    assert revoke.status_code == 200


def test_capital_call_with_wire_instructions_requires_review(client: TestClient, db_session: Session) -> None:
    auth = bootstrap_portfolio(client, db_session, email="phasecd-capcall@example.com")
    headers = csrf_headers(auth)

    document = client.post(
        "/api/documents/upload",
        data={"folder": "capital_calls"},
        files={
            "file": (
                "BlueRiver-capital-call.xlsx",
                build_capital_call_xlsx(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        headers=headers,
    )
    assert document.status_code == 200
    document_id = document.json()["id"]

    parse = client.post(f"/api/documents/{document_id}/parse", headers=headers)
    assert parse.status_code == 200

    document_detail = client.get(f"/api/documents/{document_id}")
    assert document_detail.status_code == 200
    assert document_detail.json()["tag"] == "capital_call"
    assert document_detail.json()["extraction_status"] == "needs_review"

    extraction = client.get(f"/api/documents/{document_id}/extraction")
    assert extraction.status_code == 200
    body = extraction.json()
    assert body["needs_review_count"] >= 2
    assert body["positions"][0]["ticker"] is None
    assert "Manual review required" in body["positions"][0]["notes"]
    assert body["confidence"][0]["issues"] == ["no structured positions extracted"]

    extraction_row = db_session.scalar(
        select(ExtractionResult).where(ExtractionResult.document_id == document_id)
    )
    assert extraction_row is not None
    assert extraction_row.confidence_json.startswith("[")
    reconciliation_artifact = db_session.scalar(
        select(ExtractionArtifact).where(
            ExtractionArtifact.extraction_result_id == extraction_row.id,
            ExtractionArtifact.artifact_type == "reconciliation",
        )
    )
    assert reconciliation_artifact is not None
    assert "\"wire_instructions_require_hitl\"" in reconciliation_artifact.payload_json
    assert "\"overall_confidence\"" in reconciliation_artifact.payload_json

    review = client.get(f"/api/documents/{document_id}/review")
    assert review.status_code == 200
    review_body = review.json()
    assert review_body["classification"]["doc_type"] == "CAPITAL_CALL"
    assert review_body["treasury"]["wire_bank"] == "First Private Bank"
    assert review_body["field_reviews"][0]["field"] == "wire_instructions"
    assert review_body["field_reviews"][0]["resolved"] is False

    reviewed_positions = review_body["positions"]
    reviewed_positions[0].update(
        {
            "ticker": "BLUERIVER",
            "name": "BlueRiver Capital Call",
            "quantity": 1,
            "market_value_usd": 1250000,
            "asset_class": "private_equity",
            "geo_region": "US",
            "sector": "Alternatives",
            "market_segment": "Private Equity",
            "factor_asset_class": "private_equity",
            "factor_sector": "alternatives",
            "factor_region": "us",
            "factor_market_segment": "private_equity",
        }
    )
    reviewed = client.patch(
        f"/api/documents/{document_id}/review",
        json={
            "positions": reviewed_positions,
            "treasury": {
                "wire_bank": "First Private Bank Zurich",
                "wire_account": "9988776655",
                "wire_routing": "021000021",
                "wire_reference": "BlueRiver Q2",
                "contact_email": "ops@blueriver.com",
            },
            "resolved_fields": ["wire_instructions"],
        },
        headers=headers,
    )
    assert reviewed.status_code == 200
    reviewed_body = reviewed.json()
    assert reviewed_body["needs_review_count"] == 0
    assert reviewed_body["field_reviews"][0]["resolved"] is True
    assert reviewed_body["treasury"]["wire_bank"] == "First Private Bank Zurich"
    assert reviewed_body["treasury"]["wire_account"] == "9988776655"
    assert reviewed_body["treasury"]["wire_reference"] == "BlueRiver Q2"

    document_detail = client.get(f"/api/documents/{document_id}")
    assert document_detail.status_code == 200
    assert document_detail.json()["extraction_status"] == "done"

    approve = client.post(f"/api/documents/{document_id}/approve", headers=headers)
    assert approve.status_code == 200
    assert "snapshot" in approve.json()["detail"]


def test_briefing_export_returns_503_when_pdf_backend_is_unavailable(
    client: TestClient,
    db_session: Session,
    monkeypatch,
) -> None:
    auth = bootstrap_portfolio(client, db_session, email="phasecd3@example.com")
    headers = csrf_headers(auth)
    assert client.post("/api/var/compute", headers=headers).status_code == 200
    assert client.post("/api/risk/run", headers=headers).status_code == 200

    briefing = client.post("/api/briefings/generate", headers=headers)
    assert briefing.status_code == 200
    briefing_id = briefing.json()["id"]

    def fail_export(*_args, **_kwargs):
        raise PdfExportUnavailableError("PDF export unavailable — WeasyPrint system libraries not installed")

    monkeypatch.setattr("backend.routers.briefings.export_briefing_pdf", fail_export)

    export = client.get(f"/api/briefings/{briefing_id}/export/pdf")
    assert export.status_code == status.HTTP_503_SERVICE_UNAVAILABLE
    assert export.json()["detail"] == "PDF export unavailable — WeasyPrint system libraries not installed"


def test_position_mutations_target_the_correct_duplicate_ticker_row(client: TestClient, db_session: Session) -> None:
    auth = bootstrap_portfolio(client, db_session, email="duplicates@example.com")
    csrf = auth["csrf"]

    for name, quantity, market_value in (
        ("Apple Legacy Lot", 3, 333.0),
        ("Apple New Lot", 7, 777.0),
    ):
        created = client.post(
            "/api/portfolio/positions",
            json={
                "ticker": "AAPL",
                "name": name,
                "quantity": quantity,
                "market_value_usd": market_value,
                "asset_class": "public_equity",
                "position_currency": "USD",
            },
            headers={"X-CSRF-Token": csrf},
        )
        assert created.status_code == 201

    positions_before = client.get("/api/portfolio/positions")
    assert positions_before.status_code == 200
    duplicate_rows = [
        item
        for item in positions_before.json()["items"]
        if item["ticker"] == "AAPL" and item["name"] in {"Apple Legacy Lot", "Apple New Lot"}
    ]
    assert len(duplicate_rows) == 2
    legacy_row = next(item for item in duplicate_rows if item["name"] == "Apple Legacy Lot")
    new_row = next(item for item in duplicate_rows if item["name"] == "Apple New Lot")

    update = client.patch(
        f"/api/portfolio/positions/{legacy_row['id']}",
        json={"quantity": 11},
        headers={"X-CSRF-Token": csrf},
    )
    assert update.status_code == 200

    positions_after_update = client.get("/api/portfolio/positions")
    assert positions_after_update.status_code == 200
    updated_rows = [
        item
        for item in positions_after_update.json()["items"]
        if item["ticker"] == "AAPL" and item["name"] in {"Apple Legacy Lot", "Apple New Lot"}
    ]
    assert len(updated_rows) == 2
    assert next(item for item in updated_rows if item["name"] == "Apple Legacy Lot")["quantity"] == 11.0
    assert next(item for item in updated_rows if item["name"] == "Apple New Lot")["quantity"] == new_row["quantity"]

    refreshed_new_row = next(item for item in updated_rows if item["name"] == "Apple New Lot")
    delete = client.delete(
        f"/api/portfolio/positions/{refreshed_new_row['id']}",
        headers={"X-CSRF-Token": csrf},
    )
    assert delete.status_code == 200

    positions_after_delete = client.get("/api/portfolio/positions")
    assert positions_after_delete.status_code == 200
    final_rows = [
        item
        for item in positions_after_delete.json()["items"]
        if item["ticker"] == "AAPL" and item["name"] in {"Apple Legacy Lot", "Apple New Lot"}
    ]
    assert len(final_rows) == 1
    assert final_rows[0]["name"] == "Apple Legacy Lot"
    assert final_rows[0]["quantity"] == 11.0


def test_position_factor_tag_updates_are_persisted(client: TestClient, db_session: Session) -> None:
    auth = bootstrap_portfolio(client, db_session, email="factor-tags@example.com")
    csrf = auth["csrf"]

    current_positions = client.get("/api/portfolio/positions")
    assert current_positions.status_code == 200
    target = current_positions.json()["items"][0]

    updated = client.patch(
        f"/api/portfolio/positions/{target['id']}",
        json={
            "factor_asset_class": "private_equity",
            "factor_sector": "software",
            "factor_subsector": "application_software",
            "factor_country": "US",
            "factor_region": "north_america",
            "factor_market_segment": "buyout",
            "factor_tag_source": "manual",
            "factor_tag_confidence": 1.0,
        },
        headers={"X-CSRF-Token": csrf},
    )
    assert updated.status_code == 200

    positions_after = client.get("/api/portfolio/positions")
    assert positions_after.status_code == 200
    refreshed = next(item for item in positions_after.json()["items"] if item["ticker"] == target["ticker"])
    assert refreshed["factor_asset_class"] == "private_equity"
    assert refreshed["factor_sector"] == "software"
    assert refreshed["factor_subsector"] == "application_software"
    assert refreshed["factor_country"] == "US"
    assert refreshed["factor_region"] == "north_america"
    assert refreshed["factor_market_segment"] == "buyout"
    assert refreshed["factor_tag_source"] == "manual"
    assert refreshed["factor_tag_confidence"] == 1.0


def test_cross_workspace_ids_are_hidden_across_documents_positions_and_briefings(
    client: TestClient,
    db_session: Session,
) -> None:
    auth = bootstrap_portfolio(client, db_session, email="workspace-a@example.com")
    headers = csrf_headers(auth)

    user_b = seed_user(db_session, email="workspace-b@example.com")
    snapshot_b = PortfolioSnapshot(
        workspace_id=user_b.workspace_id,
        uploaded_by=user_b.id,
        source="csv",
        position_count=1,
        total_aum_usd=2500.0,
        is_current=True,
    )
    db_session.add(snapshot_b)
    db_session.flush()
    var_result_b = VarResult(
        snapshot_id=snapshot_b.id,
        workspace_id=user_b.workspace_id,
        var_1d_95=1.0,
        var_1d_99=2.0,
        cvar_1d_95=1.5,
        cvar_1d_99=2.5,
        max_drawdown_1y=3.0,
        worst_scenario_date=utc_now().date(),
        worst_scenario_loss=100.0,
        lookback_days=252,
        effective_lookback_days=252,
        methodology="historical_simulation",
        model_coverage_pct=100.0,
        unmodeled_value_usd=0.0,
        position_contributions_json="[]",
        assumptions_json="{}",
        computed_at=utc_now(),
    )
    db_session.add(var_result_b)
    db_session.flush()

    position_b = Position(
        snapshot_id=snapshot_b.id,
        workspace_id=user_b.workspace_id,
        ticker="BETA",
        name="Beta Holding",
        position_currency="USD",
        quantity=5,
        price_usd=500.0,
        market_value_usd=2500.0,
        asset_class="public_equity",
        custodian="Other Custodian",
        price_source="manual",
    )
    document_b = Document(
        workspace_id=user_b.workspace_id,
        uploaded_by=user_b.id,
        filename="other.pdf",
        file_type="pdf",
        file_size_bytes=128,
        sha256="b" * 64,
        storage_path="/tmp/other.pdf",
        folder="custodian_statements",
        malware_scan_status="clean",
        extraction_status="pending",
    )
    briefing_b = BriefingRun(
        workspace_id=user_b.workspace_id,
        snapshot_id=snapshot_b.id,
        var_result_id=var_result_b.id,
        generated_by=user_b.id,
        version=1,
        week_label="week-15-2026",
        status="draft",
        output_json="{}",
        model="claude-opus-4-6",
        created_at=utc_now(),
    )
    db_session.add_all([position_b, document_b, briefing_b])
    db_session.commit()

    assert client.get(f"/api/documents/{document_b.id}").status_code == 404
    assert client.get(f"/api/briefings/{briefing_b.id}").status_code == 404

    position_update = client.patch(
        f"/api/portfolio/positions/{position_b.id}",
        json={"quantity": 9},
        headers=headers,
    )
    assert position_update.status_code == 404


def test_protected_routes_reject_unauthenticated_requests(client: TestClient) -> None:
    checks = [
        ("GET", "/api/auth/me", None),
        ("GET", "/api/documents", None),
        ("GET", "/api/portfolio/positions", None),
        ("POST", "/api/var/compute", {}),
        ("POST", "/api/risk/run", {}),
    ]

    for method, path, payload in checks:
        if method == "GET":
            response = client.get(path)
        else:
            response = client.post(path, json=payload)
        assert response.status_code == 401, path
